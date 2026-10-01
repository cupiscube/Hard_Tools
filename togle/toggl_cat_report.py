"""
toggl_cat_report.py
===================
Собирает ежемесячные «Отчёты по задачам» по пунктам договора (теги CAT:)
из детальной выгрузки toggl_detailed_export.py — отдельный файл на каждого юзера,
в формате клиентского шаблона (Детализация / Data_collected / Свод по задачам).

Дополнительно создаётся внутренний файл _Свод_CAT_*.xlsx:
  * «Часы по категориям» — люди × пункты договора (формулы SUMIFS)
  * «Проверка»           — записи без CAT, со старой схемой CAT, спорные
  * «Сопоставление тегов» — какой тег в какой пункт ушёл и почему
  * «Все записи»         — плоская таблица для своей аналитики / сводных

Как теги сопоставляются с пунктами (безопасность важнее полноты):
  1. Тег из списка NEW_CAT_TAGS (номер + начало названия) → пункт договора.
  2. Тег из списка LEGACY_CAT_TAGS (старая схема 1–8) → LEGACY_TO_CLAUSE,
     а если там не задано — НЕ считается, попадает в «Проверку».
  3. Неизвестный тег «CAT: N» → пункт N, но только если у человека
     нет тегов старой схемы (иначе номер может означать другое). Помечается.

Запуск:
  python toggl_cat_report.py                       # берёт свежий toggl_detailed_*.xlsx рядом
  python toggl_cat_report.py --input toggl_detailed_team_2026-09-01_2026-09-30.xlsx
  python toggl_cat_report.py --people BADA
"""

import argparse
import datetime as dt
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================================================================
# STEP 0: НАСТРОЙКИ
# ============================================================================

INPUT_FILES = None            # None = самый свежий toggl_detailed_*.xlsx рядом со скриптом
                              # или список путей: ["toggl_detailed_team_....xlsx"]
INPUT_SHEET = "Detailed"
PEOPLE = None                 # None = все, кто есть в выгрузке; или ['BADA', 'CHAL']
OUTPUT_DIR = None             # None = <папка скрипта>/CAT_reports_<YYYY-MM>
SKIP_USERS_WITHOUT_CAT = True # не создавать пустой отчёт тем, у кого нет ни одного CAT
MINUTES_DECIMALS = 2          # округление минут в записи (0 = до целых минут)
LOG_ENTRIES_LIMIT = 10        # сколько проблемных записей печатать в консоль на каждый тип
                              # предупреждения (остальные — на листе «Проверка»); 0 = не печатать

# Фамилии для имени файла: "<Фамилия>_Отчет_по_задачам_<месяц>_<год>.xlsx"
SURNAMES = {
    "BADA": "Батраков",
}
# Ставка, руб./час (ячейка H2). None = оставить пустой жёлтой для ручного ввода
RATES = {
    # "BADA": 2500,
}

# Пункты договора — строки отчёта (номер → текст в колонке B)
CONTRACT_CLAUSES = {
    1: "•  Подготовка отчётов о качестве данных из различных источников в корпоративном хранилище данных (DWH) Заказчика и клиентов Заказчика;",
    2: "•  Разработка и сопровождение автоматических проверок качества данных, в том числе с использованием каталогов данных и метаданных (OpenMetadata и аналогичных);",
    3: "•  Первичная обработка данных из различных источников с предоставлением Заказчику структурированного набора данных в согласованном формате;",
    4: "•  Связывание (интеграция) данных из различных баз данных и информационных систем;",
    5: "•  Проектирование и создание витрин данных в реляционных и аналитических СУБД (PostgreSQL, ClickHouse и аналогичных);",
    6: "•  Разработка ETL/ELT-процессов и конвейеров обработки данных, включая автоматическую сборку данных по расписанию и средства распределённой обработки данных (Apache Spark и аналогичные);",
    7: "•  Загрузка, организация и сопровождение данных в облачных хранилищах (Azure Blob Storage и аналогичных);",
    8: "•  Разработка алгоритмов автоматического сопоставления (матчинга) и классификации данных, в том числе с применением методов машинного обучения и ИИ-инструментов;",
    9: "•  Создание дашбордов и аналитической отчётности в BI-системах (Power BI, Superset и аналогичных);",
    10: "•  Написание технической документации (описания наборов данных, витрин данных, аналитических отчётов, стандартов именования);",
    11: "•  Иные услуги в области обработки и анализа данных по дополнительному запросу Заказчика.",
    12: "•  Дополнительные выплаты (премия / оплата взносов и пр.)",
}
EXTRA_CLAUSE = 12             # строка «Дополнительные выплаты»: рубли вносятся вручную
EXTRA_RUBLES_BY_RATE = False  # True = считать рубли строки 12 как часы × ставка

# Новая схема (CAT: 0–11): (номер тега, начало названия) → пункт договора.
# Номера 1, 2, 3, 8 пока не встречались — допишите их названия, когда появятся.
NEW_CAT_TAGS = [
    (0, "Дополнительные выплаты", 12),
    (4, "Связывание (интеграция) данных", 4),
    (5, "Проектирование и создание витрин данных", 5),
    (6, "Разработка ETL/ELT-процессов", 6),
    (7, "Загрузка и сопровождение данных в облачных хранилищах", 7),
    (9, "Дашборды и аналитическая отчётность в BI-системах", 9),
    (10, "Техническая документация", 10),
    (11, "Иные услуги по обработке и анализу данных", 11),
]
# Старая схема (CAT: 1–8) — номера означают ДРУГИЕ пункты
LEGACY_CAT_TAGS = [
    (1, "Исследование разрозненных источников данных"),
    (2, "Планирование сбора данных в хранилище"),
    (3, "Оценка чистоты данных"),
    (4, "Разработка и доработка BI-отчетов и дашбордов"),
    (7, "Услуги по аналитике данных в рамках консалтинговых проектов"),
]
# Перенос старых тегов на новые пункты договора. Пусто = не считать (уйдут в «Проверку»).
# Пример (ПРОВЕРЬТЕ по договору, прежде чем включать):
LEGACY_TO_CLAUSE = {
    # 1: 3,   # Исследование разрозненных источников → Первичная обработка данных
    # 2: 6,   # Планирование сбора данных (наладка ETL) → ETL/ELT-процессы
    # 3: 2,   # Оценка чистоты данных → Автоматические проверки качества
    # 4: 9,   # BI-отчёты и дашборды → Дашборды в BI-системах
    # 7: 11,  # Аналитика в консалтинговых проектах → Иные услуги
}

# ============================================================================

try:
    SCRIPT_DIR = Path(__file__).resolve().parent
except NameError:
    SCRIPT_DIR = Path.cwd()

MONTHS_NOM = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль",
              "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
MONTHS_PREP = ["январе", "феврале", "марте", "апреле", "мае", "июне", "июле",
               "августе", "сентябре", "октябре", "ноябре", "декабре"]

_HOMOGLYPHS = str.maketrans("АВЕКМНОРСТУХ", "ABEKMHOPCTYX")
# Тег CAT: номер, затем название до следующего тега вида ", XX: " или конца строки
_CAT_RE = re.compile(r"CAT:\s*0*(\d+)\s*(?:\|\s*)?(.*?)(?=,\s*[A-Za-zА-Яа-яЁё]{1,12}:\s|$)")


def log(msg):
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)


def norm_text(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).replace("ё", "е").replace("Ё", "Е").strip().casefold()


def norm_abbr(s) -> str:
    return str(s).strip().translate(_HOMOGLYPHS).upper()


# ============================================================================
# Чтение выгрузки
# ============================================================================
REQUIRED = ["Abbreviation", "Description", "Project", "Start date", "Tags"]


def find_inputs() -> list:
    if INPUT_FILES:
        return [Path(p) if Path(p).is_absolute() else SCRIPT_DIR / p for p in INPUT_FILES]
    cands = sorted(set(SCRIPT_DIR.glob("toggl_detailed_*.xlsx")) | set(Path.cwd().glob("toggl_detailed_*.xlsx")),
                   key=lambda p: p.stat().st_mtime, reverse=True)
    cands = [p for p in cands if not p.name.startswith("~$")]
    if not cands:
        raise FileNotFoundError("Не найден toggl_detailed_*.xlsx рядом со скриптом — укажите --input")
    return [cands[0]]


def duration_seconds(df: pd.DataFrame) -> pd.Series:
    if "Duration" in df.columns:
        def hms(v):
            if pd.isna(v):
                return None
            if isinstance(v, dt.time):
                return v.hour * 3600 + v.minute * 60 + v.second
            if isinstance(v, dt.timedelta):
                return v.total_seconds()
            parts = str(v).strip().split(":")
            if len(parts) == 3:
                return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
            return None
        sec = df["Duration"].map(hms)
        if sec.notna().all():
            return sec.astype(float)
    if "Duration (h)" in df.columns:
        return (pd.to_numeric(df["Duration (h)"], errors="coerce") * 3600).round()
    raise KeyError("В выгрузке нет ни 'Duration', ни 'Duration (h)'")


def load_entries(paths: list) -> pd.DataFrame:
    frames = []
    for p in paths:
        log(f"📥 Читаю {p.name}")
        if p.suffix.lower() == ".csv":
            df = pd.read_csv(p, sep=None, engine="python", encoding="utf-8-sig")
        else:
            df = pd.read_excel(p, sheet_name=INPUT_SHEET)
        missing = [c for c in REQUIRED if c not in df.columns]
        if missing:
            raise KeyError(f"{p.name}: нет обязательных колонок {missing}")
        frames.append(df)
    df = pd.concat(frames, ignore_index=True)
    if "Time entry ID" in df.columns:
        before = len(df)
        df = df.drop_duplicates(subset="Time entry ID")
        if len(df) < before:
            log(f"   ↺ убрано дублей: {before - len(df)}")

    df["Abbreviation"] = df["Abbreviation"].map(norm_abbr)
    df["Description"] = df["Description"].fillna("").astype(str)
    df["Project"] = df["Project"].fillna("").astype(str)
    df["Tags"] = df["Tags"].fillna("").astype(str)
    df["Date"] = pd.to_datetime(df["Start date"], errors="coerce", dayfirst=True)
    st = df["Start time"].astype(str) if "Start time" in df.columns else "00:00:00"
    df["_sort"] = pd.to_datetime(df["Date"].dt.strftime("%Y-%m-%d") + " " + st, errors="coerce")
    df["Seconds"] = duration_seconds(df)
    df["Minutes"] = (df["Seconds"] / 60).round(MINUTES_DECIMALS)
    if MINUTES_DECIMALS == 0:
        df["Minutes"] = df["Minutes"].astype(int)
    return df.sort_values(["Abbreviation", "_sort"]).reset_index(drop=True)


# ============================================================================
# Сопоставление тегов → пункт договора
# ============================================================================
_NEW = [(n, norm_text(prefix), clause) for n, prefix, clause in NEW_CAT_TAGS]
_LEGACY = [(n, norm_text(prefix)) for n, prefix in LEGACY_CAT_TAGS]


def classify_tag(num: int, name: str):
    """→ (scheme, clause|None).  scheme: new | legacy | unknown"""
    nn = norm_text(name)
    for n, prefix, clause in _NEW:
        if n == num and nn.startswith(prefix):
            return "new", clause
    for n, prefix in _LEGACY:
        if n == num and nn.startswith(prefix):
            return "legacy", LEGACY_TO_CLAUSE.get(num)
    return "unknown", (num if num in CONTRACT_CLAUSES and num != EXTRA_CLAUSE else None)


def assign_clauses(df: pd.DataFrame) -> pd.DataFrame:
    cats = df["Tags"].map(lambda s: [(int(n), name.strip()) for n, name in _CAT_RE.findall(s)])
    parsed = cats.map(lambda lst: [(n, name, *classify_tag(n, name)) for n, name in lst])
    df["_cats"] = parsed
    users_with_legacy = set(df.loc[parsed.map(lambda l: any(x[2] == "legacy" for x in l)), "Abbreviation"])

    clause, status, cat_tag = [], [], []
    for abbr, lst in zip(df["Abbreviation"], parsed):
        if not lst:
            clause.append(None); status.append("Нет тега CAT"); cat_tag.append("")
            continue
        n, name, scheme, cl = lst[0]
        tag = f"CAT: {n} | {name}"
        note = " (+ ещё CAT-теги: учтён первый)" if len(lst) > 1 else ""
        if scheme == "new":
            clause.append(cl); status.append("OK" + note)
        elif scheme == "legacy" and cl is not None:
            clause.append(cl); status.append(f"Старая схема → пункт {cl} (LEGACY_TO_CLAUSE)" + note)
        elif scheme == "legacy":
            clause.append(None); status.append("Старая схема CAT — не учтено (задайте LEGACY_TO_CLAUSE)")
        elif cl is not None and abbr not in users_with_legacy:
            clause.append(cl); status.append("Сопоставлено только по номеру — проверьте название" + note)
        else:
            clause.append(None); status.append("Неизвестный CAT-тег — не учтено")
        cat_tag.append(tag)
    df["Clause"] = clause
    df["Mapping"] = status
    df["CAT tag"] = cat_tag
    return df


# ============================================================================
# Стили (как в клиентском шаблоне)
# ============================================================================
FONT = "Arial"
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="FFDDEBF7")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFFFFF00")
BLUE = "FF0000FF"
NF_2 = "0.00"
NF_RUB = "#\\ ##0"


def f(bold=False, size=11, italic=False, color=None):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def put(ws, ref, value=None, font=None, fill=None, border=None, nf=None, align=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font or f()
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if nf:
        c.number_format = nf
    if align:
        c.alignment = align
    return c


def header_row(ws, row, headers, align_center=True, wrap=False):
    for i, h in enumerate(headers, 1):
        put(ws, f"{get_column_letter(i)}{row}", h, font=f(bold=True), fill=HEAD_FILL, border=BOX,
            align=Alignment(horizontal="center" if align_center else None, wrap_text=wrap or None))


# ============================================================================
# Клиентский отчёт одного человека
# ============================================================================
def period_info(dates: pd.Series):
    d0, d1 = dates.min(), dates.max()
    if d0.year == d1.year and d0.month == d1.month:
        return {"same_month": True, "year": d0.year, "nom": MONTHS_NOM[d0.month - 1],
                "prep": MONTHS_PREP[d0.month - 1], "ym": f"{d0:%Y-%m}"}
    return {"same_month": False, "year": d0.year, "ym": f"{d0:%Y-%m-%d}_{d1:%Y-%m-%d}",
            "range": f"{d0:%d.%m.%Y}–{d1:%d.%m.%Y}"}


def zero_note(empty: list, per: dict) -> str:
    if not empty:
        return ""
    when = f"в {per['prep']} {per['year']} г." if per["same_month"] else f"за период {per['range']}"
    if len(empty) == 1:
        return f"Категория {empty[0]} {when} работ не содержит."
    nums = ", ".join(map(str, empty[:-1])) + f" и {empty[-1]}"
    return f"Категории {nums} {when} работ не содержат."


def build_user_report(abbr: str, data: pd.DataFrame, per: dict, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Детализация за месяц {per['year']}"
    dc = wb.create_sheet("Data_collected")
    sv = wb.create_sheet("Свод по задачам")

    # ---------- Data_collected ----------
    header_row(dc, 1, ["Description", "category", "Duration", "Project", "Date"], align_center=False)
    for i, r in enumerate(data.itertuples(index=False), start=2):
        put(dc, f"A{i}", r.Description, border=BOX)
        put(dc, f"B{i}", int(r.Clause), border=BOX, align=Alignment(horizontal="center"))
        put(dc, f"C{i}", r.Minutes, border=BOX, nf=NF_2)
        put(dc, f"D{i}", r.Project, border=BOX)
        put(dc, f"E{i}", r.Date.to_pydatetime() if pd.notna(r.Date) else None, border=BOX, nf="DD.MM.YYYY")
    last = max(len(data) + 1, 2)
    for col, w in zip("ABCDE", [62, 10, 11, 22, 12]):
        dc.column_dimensions[col].width = w
    dc.freeze_panes = "A2"

    # ---------- Детализация ----------
    title = (f"Отчет по задачам за {per['nom']} {per['year']}" if per["same_month"]
             else f"Отчет по задачам за период {per['range']}")
    ws.merge_cells("A1:B1")
    put(ws, "A1", title, font=f(bold=True, size=12))
    for col, h in zip("CDE", ["минуты", "часы", "рубли"]):
        put(ws, f"{col}2", h, font=f(bold=True), fill=HEAD_FILL, border=BOX,
            align=Alignment(horizontal="center"))
    put(ws, "G2", "Ставка, руб./час:", font=f(bold=True))
    put(ws, "H2", RATES.get(abbr), font=f(color=BLUE), fill=INPUT_FILL, border=BOX, nf=NF_RUB)

    rng_b = f"Data_collected!$B$2:$B${last}"
    rng_c = f"Data_collected!$C$2:$C${last}"
    for idx, (num, text) in enumerate(CONTRACT_CLAUSES.items()):
        r = 3 + idx
        put(ws, f"A{r}", num, align=Alignment(horizontal="center", vertical="top"))
        put(ws, f"B{r}", text, align=Alignment(vertical="top", wrap_text=True))
        put(ws, f"C{r}", f"=SUMIF({rng_b},$A{r},{rng_c})", border=BOX, nf=NF_2)
        put(ws, f"D{r}", f"=C{r}/60", border=BOX, nf=NF_2)
        if num == EXTRA_CLAUSE and not EXTRA_RUBLES_BY_RATE:
            put(ws, f"E{r}", None, font=f(color=BLUE), fill=INPUT_FILL, border=BOX, nf=NF_RUB)
        else:
            put(ws, f"E{r}", f"=ROUND(D{r}*$H$2,0)", border=BOX, nf=NF_RUB)
        if len(text) > 110:
            ws.row_dimensions[r].height = 30
    tot = 3 + len(CONTRACT_CLAUSES)
    first, lastc = 3, tot - 1
    put(ws, f"B{tot}", "Итого:", font=f(bold=True), align=Alignment(horizontal="right"))
    for col in "CDE":
        put(ws, f"{col}{tot}", f"=SUM({col}{first}:{col}{lastc})", font=f(bold=True),
            fill=TOTAL_FILL, border=BOX, nf=NF_RUB if col == "E" else NF_2)

    notes = ["Жёлтые ячейки заполняются вручную. Столбец «рубли» = часы × ставка."]
    used = set(data["Clause"].astype(int))
    empty = [n for n in CONTRACT_CLAUSES if n != EXTRA_CLAUSE and n not in used]
    if zero_note(empty, per):
        notes.append(zero_note(empty, per))
    if EXTRA_CLAUSE in used and not EXTRA_RUBLES_BY_RATE:
        notes.append(f"Пункт {EXTRA_CLAUSE}: время по тегу CAT: 0, сумма в рублях вносится вручную.")
    for i, note in enumerate(notes):
        put(ws, f"G{3 + i}", note, font=f(size=9, italic=True, color="FF808080"))

    # список записей
    hdr = tot + 4
    header_row(ws, hdr, ["#", "Description", "category#", "Duration", "Project"])
    for i, r in enumerate(data.itertuples(index=False), start=1):
        row = hdr + i
        put(ws, f"A{row}", i, border=BOX, align=Alignment(horizontal="center"))
        put(ws, f"B{row}", r.Description, border=BOX)
        put(ws, f"C{row}", int(r.Clause), border=BOX, align=Alignment(horizontal="center"))
        put(ws, f"D{row}", r.Minutes, border=BOX, nf=NF_2)
        put(ws, f"E{row}", r.Project, border=BOX)
    end = hdr + len(data) + 1
    put(ws, f"B{end}", "Итого минут:", font=f(bold=True), align=Alignment(horizontal="right"))
    put(ws, f"D{end}", f"=SUM(D{hdr + 1}:D{end - 1})", font=f(bold=True), fill=TOTAL_FILL,
        border=BOX, nf=NF_2)
    for col, w in {"A": 5.5, "B": 100, "C": 12, "D": 11, "E": 25, "G": 22, "H": 14}.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 16
    ws.freeze_panes = "A3"
    ws.sheet_view.zoomScale = 116

    # ---------- Свод по задачам ----------
    header_row(sv, 1, ["Пункт договора", "Задача (Toggl)", "Записей", "Минуты", "Часы"], wrap=True)
    sv.row_dimensions[1].height = 31
    grp = (data.groupby(["Clause", "Description"], sort=False)
           .agg(n=("Minutes", "size"), m=("Minutes", "sum")).reset_index()
           .sort_values(["Clause", "m"], ascending=[True, False]))
    for i, r in enumerate(grp.itertuples(index=False), start=2):
        put(sv, f"A{i}", int(r.Clause), border=BOX, align=Alignment(horizontal="center"))
        put(sv, f"B{i}", r.Description, border=BOX)
        put(sv, f"C{i}", int(r.n), border=BOX, align=Alignment(horizontal="center"))
        put(sv, f"D{i}", round(float(r.m), MINUTES_DECIMALS), border=BOX, nf=NF_2)
        put(sv, f"E{i}", f"=D{i}/60", border=BOX, nf=NF_2)
    t = len(grp) + 2
    put(sv, f"B{t}", "Итого:", font=f(bold=True), align=Alignment(horizontal="right"))
    for col, nf in zip("CDE", ["0", NF_2, NF_2]):
        put(sv, f"{col}{t}", f"=SUM({col}2:{col}{t - 1})", font=f(bold=True), fill=TOTAL_FILL,
            border=BOX, nf=nf)
    for col, w in zip("ABCDE", [14, 64, 10, 11, 10]):
        sv.column_dimensions[col].width = w
    sv.freeze_panes = "A2"

    wb.save(path)


# ============================================================================
# Внутренний сводный файл (аналитика + проверка)
# ============================================================================
def build_team_file(df: pd.DataFrame, per: dict, path: Path, generated: dict):
    wb = Workbook()
    hs = wb.active
    hs.title = "Часы по категориям"
    allr = wb.create_sheet("Все записи")
    chk = wb.create_sheet("Проверка")
    mp = wb.create_sheet("Сопоставление тегов")

    # --- Все записи ---
    cols = ["Abbreviation", "User", "Date", "Description", "Project", "Minutes", "Hours",
            "Clause", "Mapping", "CAT tag", "Tags", "Time entry ID"]
    flat = df.assign(
        User=df["User"] if "User" in df.columns else "",
        Hours=df["Minutes"] / 60,
        Clause=df["Clause"].map(lambda c: int(c) if pd.notna(c) else "нет"),
        **({"Time entry ID": ""} if "Time entry ID" not in df.columns else {}),
    )[cols]
    header_row(allr, 1, cols, align_center=False)
    for i, row in enumerate(flat.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            c = allr.cell(i, j, v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
            c.font = f()
            if cols[j - 1] == "Date":
                c.number_format = "DD.MM.YYYY"
            elif cols[j - 1] in ("Minutes", "Hours"):
                c.number_format = NF_2
    n_last = max(len(flat) + 1, 2)
    allr.freeze_panes = "A2"
    allr.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n_last}"
    for j, w in enumerate([11, 22, 11, 60, 26, 9, 8, 8, 45, 50, 60, 14], start=1):
        allr.column_dimensions[get_column_letter(j)].width = w

    # --- Часы по категориям (SUMIFS по «Все записи») ---
    put(hs, "A1", f"Часы по пунктам договора — {per.get('nom', '')} {per['year']}".strip(),
        font=f(bold=True, size=12))
    heads = ["Юзер", "Отчёт создан"] + [str(n) for n in CONTRACT_CLAUSES] + \
            ["Учтено, ч", "Без CAT / не учтено, ч", "Всего в Toggl, ч", "Доля учтённого"]
    header_row(hs, 3, heads, wrap=True)
    hs.row_dimensions[3].height = 31
    A, H, CL = "'Все записи'!$A$2:$A$%d" % n_last, "'Все записи'!$G$2:$G$%d" % n_last, \
        "'Все записи'!$H$2:$H$%d" % n_last
    users = sorted(df["Abbreviation"].unique())
    ncl = len(CONTRACT_CLAUSES)
    for i, u in enumerate(users, start=4):
        put(hs, f"A{i}", u, font=f(bold=True), border=BOX)
        put(hs, f"B{i}", "да" if generated.get(u) else "нет", border=BOX,
            align=Alignment(horizontal="center"))
        for k, num in enumerate(CONTRACT_CLAUSES, start=3):
            put(hs, f"{get_column_letter(k)}{i}", f"=SUMIFS({H},{A},$A{i},{CL},{num})",
                border=BOX, nf=NF_2)
        c_first, c_last = get_column_letter(3), get_column_letter(2 + ncl)
        cu, cn, ct, cs = (get_column_letter(3 + ncl + x) for x in range(4))
        put(hs, f"{cu}{i}", f"=SUM({c_first}{i}:{c_last}{i})", font=f(bold=True), border=BOX, nf=NF_2)
        put(hs, f"{cn}{i}", f'=SUMIFS({H},{A},$A{i},{CL},"нет")', border=BOX, nf=NF_2)
        put(hs, f"{ct}{i}", f"=SUMIFS({H},{A},$A{i})", border=BOX, nf=NF_2)
        put(hs, f"{cs}{i}", f"=IF({ct}{i}=0,0,{cu}{i}/{ct}{i})", border=BOX, nf="0%")
    tr = 4 + len(users)
    put(hs, f"A{tr}", "Итого", font=f(bold=True), fill=TOTAL_FILL, border=BOX)
    put(hs, f"B{tr}", None, fill=TOTAL_FILL, border=BOX)
    for k in range(3, 3 + ncl + 3):
        col = get_column_letter(k)
        put(hs, f"{col}{tr}", f"=SUM({col}4:{col}{tr - 1})", font=f(bold=True), fill=TOTAL_FILL,
            border=BOX, nf=NF_2)
    cu, ct, cs = get_column_letter(3 + ncl), get_column_letter(5 + ncl), get_column_letter(6 + ncl)
    put(hs, f"{cs}{tr}", f"=IF({ct}{tr}=0,0,{cu}{tr}/{ct}{tr})", font=f(bold=True), fill=TOTAL_FILL,
        border=BOX, nf="0%")
    hs.column_dimensions["A"].width = 10
    hs.column_dimensions["B"].width = 10
    for k in range(3, 3 + ncl):
        hs.column_dimensions[get_column_letter(k)].width = 8
    for k in range(3 + ncl, 7 + ncl):
        hs.column_dimensions[get_column_letter(k)].width = 13
    hs.freeze_panes = "C4"
    put(hs, f"A{tr + 2}", "Номера колонок = пункты договора (лист «Сопоставление тегов» / CONTRACT_CLAUSES). "
        "Формулы ссылаются на лист «Все записи».", font=f(size=9, italic=True, color="FF808080"))

    # --- Проверка ---
    bad = df[df["Mapping"] != "OK"].copy()
    for opt in ("Start time", "Time entry ID"):
        if opt not in bad.columns:
            bad[opt] = ""
    ccols = ["Abbreviation", "Date", "Start time", "Description", "Project", "Minutes",
             "Mapping", "Tags", "Time entry ID"]
    header_row(chk, 1, ccols, align_center=False)
    for i, row in enumerate(bad[ccols].itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            c = chk.cell(i, j, v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
            c.font = f()
            if ccols[j - 1] == "Date":
                c.number_format = "DD.MM.YYYY"
            elif ccols[j - 1] == "Minutes":
                c.number_format = NF_2
    chk.freeze_panes = "A2"
    chk.auto_filter.ref = f"A1:I{max(len(bad) + 1, 2)}"
    for j, w in enumerate([11, 11, 10, 60, 26, 9, 55, 70, 14], start=1):
        chk.column_dimensions[get_column_letter(j)].width = w

    # --- Сопоставление тегов ---
    rows = {}
    for abbr, lst, mins in zip(df["Abbreviation"], df["_cats"], df["Minutes"]):
        for n, name, scheme, cl in lst:
            key = (abbr, n, name)
            r = rows.setdefault(key, {"scheme": scheme, "clause": cl, "n": 0, "min": 0.0})
            r["n"] += 1
            r["min"] += mins
    mcols = ["Abbreviation", "CAT", "Название тега", "Схема", "→ Пункт договора", "Записей", "Часы"]
    header_row(mp, 1, mcols, align_center=False)
    scheme_ru = {"new": "новая", "legacy": "СТАРАЯ", "unknown": "неизвестная (по номеру)"}
    for i, ((abbr, n, name), r) in enumerate(sorted(rows.items()), start=2):
        vals = [abbr, n, name, scheme_ru[r["scheme"]], r["clause"] if r["clause"] is not None else "не учтено",
                r["n"], round(r["min"] / 60, 2)]
        for j, v in enumerate(vals, start=1):
            mp.cell(i, j, v).font = f(bold=(j == 4 and r["scheme"] != "new"))
    for j, w in enumerate([11, 6, 80, 22, 18, 9, 8], start=1):
        mp.column_dimensions[get_column_letter(j)].width = w
    mp.freeze_panes = "A2"

    wb.save(path)


def print_problem_entries(rows: pd.DataFrame) -> None:
    """Печатает записи так, чтобы их можно было найти в Toggl:
    дата, время начала, длительность, проект, описание, CAT-тег, ID записи."""
    if LOG_ENTRIES_LIMIT <= 0:
        return
    for _, r in rows.head(LOG_ENTRIES_LIMIT).iterrows():
        if pd.notna(r.get("_sort")):
            when = r["_sort"].strftime("%d.%m.%Y %H:%M")
        elif pd.notna(r.get("Date")):
            when = r["Date"].strftime("%d.%m.%Y")
        else:
            when = "?"
        mins = int(round(r["Seconds"] / 60)) if pd.notna(r.get("Seconds")) else 0
        line = f"      • {when} ({mins // 60}:{mins % 60:02d} ч) | {r['Project']} | {r['Description']}"
        if r.get("CAT tag"):
            line += f" | тег: {r['CAT tag']}"
        eid = r.get("Time entry ID")
        if eid is not None and pd.notna(eid) and str(eid).strip():
            line += f" | ID {int(eid)}"
        log(line)
    if len(rows) > LOG_ENTRIES_LIMIT:
        log(f"      … и ещё {len(rows) - LOG_ENTRIES_LIMIT} — полный список на листе «Проверка»")


# ============================================================================
def main():
    global PEOPLE, OUTPUT_DIR
    ap = argparse.ArgumentParser(description="Отчёты по пунктам договора (CAT) из выгрузки Toggl")
    ap.add_argument("--input", nargs="+", help="файл(ы) toggl_detailed_*.xlsx / .csv")
    ap.add_argument("--people", nargs="+")
    ap.add_argument("--out")
    args, _ = ap.parse_known_args()

    paths = [Path(p) for p in args.input] if args.input else find_inputs()
    df = load_entries(paths)
    df = df[df["Date"].notna()]
    if args.people or PEOPLE:
        keep = {norm_abbr(x) for x in (args.people or PEOPLE)}
        df = df[df["Abbreviation"].isin(keep)]
    if df.empty:
        raise SystemExit("❌ Нет записей для обработки")
    df = assign_clauses(df).reset_index(drop=True)

    per = period_info(df["Date"])
    out_dir = Path(args.out or OUTPUT_DIR or SCRIPT_DIR / f"CAT_reports_{per['ym']}")
    out_dir.mkdir(parents=True, exist_ok=True)
    log(f"📅 Период: {per.get('nom', per.get('range'))} {per['year']} | записей: {len(df)} "
        f"| людей: {df['Abbreviation'].nunique()}")

    generated = {}
    for abbr, g in df.groupby("Abbreviation"):
        data = g[g["Clause"].notna()].copy()
        stats = g["Mapping"].value_counts().to_dict()
        total_h, used_h = g["Minutes"].sum() / 60, data["Minutes"].sum() / 60
        log(f"##### {abbr}: всего {total_h:.2f} ч, по пунктам договора {used_h:.2f} ч")
        for k, v in stats.items():
            if k != "OK":
                log(f"   ⚠️ {k}: {v} зап.")
                print_problem_entries(g[g["Mapping"] == k])
        if data.empty and SKIP_USERS_WITHOUT_CAT:
            log("   ⏭  нет учитываемых CAT-записей — отчёт не создан")
            generated[abbr] = False
            continue
        surname = SURNAMES.get(abbr, abbr)
        name = (f"{surname}_Отчет_по_задачам_{per['nom']}_{per['year']}.xlsx" if per["same_month"]
                else f"{surname}_Отчет_по_задачам_{per['ym']}.xlsx")
        build_user_report(abbr, data, per, out_dir / name)
        generated[abbr] = True
        log(f"   💾 {name}")

    team = out_dir / f"_Свод_CAT_{per['ym']}.xlsx"
    build_team_file(df, per, team, generated)
    log(f"💾 {team.name}  (часы по категориям + лист «Проверка»)")
    log(f"📂 {out_dir}")


if __name__ == "__main__":
    main()