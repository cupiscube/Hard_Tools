"""
toggl_detailed_export.py
========================
Выгрузка детального отчёта Toggl Track (аналог Reports → Detailed → Export)
за произвольный период в Excel и/или CSV — по одному человеку или по многим.

Файлы сохраняются рядом с этим скриптом (или в OUTPUT_DIR, если задан).

Откуда берутся данные:
  * "timeline" — API v9  GET /me/time_entries  (1 запрос на ~1000 записей,
                 но Toggl отдаёт данные только за последние ~3 месяца)
  * "reports"  — Reports API v3  POST .../search/time_entries
                 (любой период, но пагинация по 50 записей = больше запросов)
  * "auto"     — timeline, если период свежий; иначе reports (рекомендуется)

Лимиты Toggl (Free): 30 запросов/час на пользователя, отдельно для /me-запросов
и для workspace-запросов. Скрипт считает запросы сам и, если упёрся в лимит
(HTTP 402/429), ждёт ровно столько, сколько говорит Toggl.

Запуск:
  python toggl_detailed_export.py                          # настройки из STEP 0
  python toggl_detailed_export.py --month 2026-09
  python toggl_detailed_export.py --start 2026-09-01 --end 2026-09-30 --people BADA SMVE
  python toggl_detailed_export.py --month 2026-09 --format both --mode file_per_person
"""

import argparse
import datetime as dt
import re
import sys
import time
from collections import deque
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import requests

# ============================================================================
# STEP 0: НАСТРОЙКИ (всё, что можно менять, — здесь)
# ============================================================================

START_DATE = "2026-09-01"   # включительно, YYYY-MM-DD
END_DATE = "2026-09-30"     # включительно, YYYY-MM-DD

# Список людей. None = все, у кого есть токен в файле токенов.
# Для одного человека: PEOPLE = ['BADA']
PEOPLE = [
    'BADA',
    # 'AIVI',
    # 'BODM',
    'CHAL',
    # 'DIDA',
    'KURO',
    # 'AMAN',
    # 'NIRO',
    # 'PRAR',
    # 'SMVE',
    # 'TODA',
    # 'VAEL',
    # 'ANAS',
    # 'NIAL',
    # 'KAAN',
    # 'LAOK',
    # 'GOBE',
]

OUTPUT_FORMAT = "xlsx"        # "xlsx" | "csv" | "both"
OUTPUT_MODE = "one_file"      # "one_file" — все в одном файле
                              # "file_per_person" — отдельный файл на каждого
SHEET_PER_PERSON = False      # one_file + xlsx: добавить лист на каждого человека
OUTPUT_DIR = None             # None = папка, где лежит скрипт
CSV_SEP = ","                 # для русской локали Excel удобнее ";"

SOURCE = "auto"               # "auto" | "timeline" | "reports"
INCLUDE_RUNNING = False       # включать ли незавершённые (идущие сейчас) записи
TIMEZONE_OVERRIDE = None      # None = часовой пояс из профиля каждого юзера,
                              # или, например, "Asia/Tbilisi" — для всех одинаково

# Файл с токенами (как в toggle_update.py). Берётся первый существующий путь.
TOKENS_FILE_CANDIDATES = [
    "./togle/data/API Tokens.xlsx",       # запуск из родительской папки
    "{script_dir}/data/API Tokens.xlsx",  # запуск из папки togle
    "{script_dir}/API Tokens.xlsx",
]
TOKENS_SHEET = "API Token"
TOKENS_ABBR_COL = "Abbreviation"
TOKENS_TOKEN_COL = "API Token"

REQUESTS_PER_HOUR = 30        # лимит Free-плана (Starter: 240, Premium: 600)
SAFETY_MARGIN = 2             # запас, чтобы не упираться в лимит впритык
TIMELINE_MAX_DAYS_BACK = 88   # Toggl: /me/time_entries не глубже ~90 дней
TIMELINE_PAGE_LIMIT = 1000    # столько максимум отдаёт /me/time_entries за раз
REPORTS_PAGE_SIZE = 50
MAX_RETRIES = 5
HTTP_TIMEOUT = 60

# ============================================================================

try:
    SCRIPT_DIR = Path(__file__).resolve().parent
except NameError:  # Jupyter
    SCRIPT_DIR = Path.cwd()

API_V9 = "https://api.track.toggl.com/api/v9"
REPORTS_V3 = "https://api.track.toggl.com/reports/api/v3"

# Кириллические двойники латинских букв (частая причина «не нашёлся токен»)
_HOMOGLYPHS = str.maketrans("АВЕКМНОРСТУХаеорсухі", "ABEKMHOPCTYXaeopcyxi")


def norm_abbr(value) -> str:
    return str(value).strip().translate(_HOMOGLYPHS).upper()


def log(msg: str) -> None:
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)


class TogglError(Exception):
    pass


# ============================================================================
# Ограничитель запросов (скользящее окно 1 час, как у Toggl)
# ============================================================================
class RateLimiter:
    def __init__(self, name: str, per_hour: int = REQUESTS_PER_HOUR,
                 margin: int = SAFETY_MARGIN):
        self.name = name
        self.limit = max(1, per_hour - margin)
        self.calls = deque()
        self.blocked_until = 0.0

    def _sleep_until(self, moment: float, reason: str) -> None:
        wait = max(0.0, moment - time.monotonic())
        if wait <= 0:
            return
        wake = dt.datetime.now() + dt.timedelta(seconds=wait)
        log(f"   😴 {self.name}: {reason}. Жду {wait / 60:.1f} мин (до {wake:%H:%M:%S})")
        time.sleep(wait)

    def acquire(self) -> None:
        if self.blocked_until > time.monotonic():
            self._sleep_until(self.blocked_until, "Toggl сообщил, что квота исчерпана")
            self.calls.clear()
            self.blocked_until = 0.0
        now = time.monotonic()
        while self.calls and now - self.calls[0] >= 3600:
            self.calls.popleft()
        if len(self.calls) >= self.limit:
            self._sleep_until(self.calls[0] + 3600 + 5, "локальный лимит запросов за час")
            return self.acquire()
        self.calls.append(time.monotonic())

    def block_for(self, seconds: float) -> None:
        self.blocked_until = max(self.blocked_until, time.monotonic() + seconds)


# ============================================================================
# Клиент Toggl для одного пользователя
# ============================================================================
class TogglUser:
    def __init__(self, abbr: str, token: str):
        self.abbr = abbr
        self.session = requests.Session()
        self.session.auth = (token, "api_token")
        self.session.headers["Content-Type"] = "application/json"
        # У Toggl два независимых лимита: /me-запросы и workspace-запросы
        self.user_bucket = RateLimiter(f"{abbr} /me")
        self.ws_bucket = RateLimiter(f"{abbr} workspace")
        self.me = None
        self.requests_made = 0
        self._projects = None           # {project_id: {...}}
        self._clients = {}              # {workspace_id: {client_id: name}}
        self._tags = {}                 # {workspace_id: {tag_id: name}}

    # ---------------- HTTP ----------------
    @staticmethod
    def _quota_wait_seconds(resp) -> float:
        for header in ("X-Toggl-Quota-Resets-In", "Retry-After"):
            val = resp.headers.get(header)
            if val and str(val).strip().isdigit():
                return int(val) + 5
        m = re.search(r"reset in (\d+) seconds", resp.text or "")
        if m:
            return int(m.group(1)) + 5
        return 600 if resp.status_code == 402 else 60

    @staticmethod
    def _is_quota_402(resp) -> bool:
        if "X-Toggl-Quota-Resets-In" in resp.headers:
            return True
        text = (resp.text or "").lower()
        return "quota" in text or "hourly limit" in text or "limit for api" in text

    def _request(self, method: str, url: str, bucket: RateLimiter, **kwargs):
        for attempt in range(1, MAX_RETRIES + 1):
            bucket.acquire()
            self.requests_made += 1
            try:
                resp = self.session.request(method, url, timeout=HTTP_TIMEOUT, **kwargs)
            except requests.RequestException as e:
                log(f"   ⚠️ Сетевая ошибка ({e}); попытка {attempt}/{MAX_RETRIES}")
                time.sleep(10 * attempt)
                continue

            remaining = resp.headers.get("X-Toggl-Quota-Remaining")
            if remaining is not None and str(remaining).strip().lstrip("-").isdigit() \
                    and int(remaining) <= 0:
                bucket.block_for(self._quota_wait_seconds(resp))

            if resp.status_code == 429 or (resp.status_code == 402 and self._is_quota_402(resp)):
                bucket.block_for(self._quota_wait_seconds(resp))
                log(f"   ⛔ HTTP {resp.status_code}: лимит API. Попытка {attempt}/{MAX_RETRIES}")
                continue
            if resp.status_code >= 500:
                log(f"   ⚠️ HTTP {resp.status_code} от Toggl; попытка {attempt}/{MAX_RETRIES}")
                time.sleep(10 * attempt)
                continue
            return resp
        raise TogglError(f"Не удалось выполнить {method} {url} за {MAX_RETRIES} попыток")

    def _get_json(self, url, bucket, params=None, soft_params=False):
        resp = self._request("GET", url, bucket, params=params)
        if resp.status_code == 400 and soft_params and params:
            resp = self._request("GET", url, bucket)  # повтор без доп. параметров
        if resp.status_code == 403:
            raise TogglError("403 Forbidden — неверный API токен или нет доступа")
        if resp.status_code != 200:
            raise TogglError(f"HTTP {resp.status_code} на {url}: {resp.text[:300]}")
        return resp.json()

    # ---------------- Справочники ----------------
    def load_me(self):
        self.me = self._get_json(f"{API_V9}/me", self.user_bucket)
        return self.me

    def projects(self) -> dict:
        if self._projects is None:
            data = self._get_json(f"{API_V9}/me/projects", self.user_bucket,
                                  params={"include_archived": "true"}, soft_params=True)
            self._projects = {p["id"]: p for p in (data or [])}
        return self._projects

    def clients(self, wid) -> dict:
        if wid not in self._clients:
            data = self._get_json(f"{API_V9}/workspaces/{wid}/clients", self.ws_bucket,
                                  params={"status": "both"}, soft_params=True)
            self._clients[wid] = {c["id"]: c.get("name") for c in (data or [])}
        return self._clients[wid]

    def tags(self, wid) -> dict:
        if wid not in self._tags:
            data = self._get_json(f"{API_V9}/workspaces/{wid}/tags", self.ws_bucket)
            self._tags[wid] = {t["id"]: t.get("name") for t in (data or [])}
        return self._tags[wid]

    # ---------------- Записи времени ----------------
    def timeline_entries(self, start_utc: dt.datetime, end_utc: dt.datetime) -> list:
        """GET /me/time_entries; если упёрлись в 1000 записей — делим период пополам."""
        params = {
            "start_date": start_utc.strftime("%Y-%m-%dT%H:%M:%SZ"),
            "end_date": end_utc.strftime("%Y-%m-%dT%H:%M:%SZ"),
            "meta": "true",
        }
        resp = self._request("GET", f"{API_V9}/me/time_entries", self.user_bucket, params=params)
        if resp.status_code == 400 and "earlier" in resp.text.lower():
            raise TimelineTooOld(resp.text)
        if resp.status_code != 200:
            raise TogglError(f"HTTP {resp.status_code} /me/time_entries: {resp.text[:300]}")
        data = resp.json() or []
        if len(data) >= TIMELINE_PAGE_LIMIT and (end_utc - start_utc) > dt.timedelta(hours=1):
            mid = start_utc + (end_utc - start_utc) / 2
            log(f"   ↔️ {len(data)} записей — делю период пополам")
            return self.timeline_entries(start_utc, mid) + self.timeline_entries(mid, end_utc)
        return data

    def reports_entries(self, wid, uid, start_date: dt.date, end_date: dt.date) -> list:
        """Reports API v3, детальный отчёт с пагинацией (возвращает список строк)."""
        url = f"{REPORTS_V3}/workspace/{wid}/search/time_entries"
        body = {
            "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(),
            "user_ids": [uid],
            "grouped": False,
            "page_size": REPORTS_PAGE_SIZE,
            "enrich_response": True,
        }
        rows, page = [], 0
        while True:
            page += 1
            resp = self._request("POST", url, self.ws_bucket, json=body)
            if resp.status_code != 200:
                raise TogglError(f"HTTP {resp.status_code} Reports API: {resp.text[:300]}")
            chunk = resp.json() or []
            rows.extend(chunk)
            next_row = resp.headers.get("X-Next-Row-Number")
            next_id = resp.headers.get("X-Next-ID")
            log(f"   📄 страница {page}: {len(chunk)} строк")
            if not chunk or not next_row:
                break
            body["first_row_number"] = int(next_row)
            if next_id:
                body["first_id"] = int(next_id)
        return rows


class TimelineTooOld(TogglError):
    pass


# ============================================================================
# Нормализация записей в единый формат
# ============================================================================
def _from_timeline(e: dict) -> dict:
    return {
        "id": e.get("id"), "workspace_id": e.get("workspace_id"),
        "project_id": e.get("project_id"), "task_id": e.get("task_id"),
        "project_name": e.get("project_name"), "client_name": e.get("client_name"),
        "task_name": e.get("task_name"),
        "description": e.get("description") or "", "billable": e.get("billable"),
        "tag_names": e.get("tags"), "tag_ids": e.get("tag_ids") or [],
        "start": e.get("start"), "stop": e.get("stop"), "seconds": e.get("duration"),
    }


def _from_reports(rows: list, wid) -> list:
    out = []
    for r in rows:
        for te in r.get("time_entries") or []:
            out.append({
                "id": te.get("id"), "workspace_id": wid,
                "project_id": r.get("project_id"), "task_id": r.get("task_id"),
                "project_name": r.get("project_name"), "client_name": r.get("client_name"),
                "task_name": r.get("task_name"),
                "description": r.get("description") or "", "billable": r.get("billable"),
                "tag_names": r.get("tag_names"), "tag_ids": r.get("tag_ids") or [],
                "start": te.get("start"), "stop": te.get("stop"), "seconds": te.get("seconds"),
            })
    return out


def resolve_names(user: TogglUser, entries: list) -> None:
    """Дозаполняет названия проектов/клиентов/тегов, если Toggl их не вернул.
    Справочники запрашиваются лениво — только если реально нужны."""
    for e in entries:
        pid, wid = e["project_id"], e["workspace_id"]
        if pid and (not e["project_name"] or not e["client_name"]):
            proj = user.projects().get(pid)
            if proj:
                e["project_name"] = e["project_name"] or proj.get("name")
                cid = proj.get("client_id")
                if not e["client_name"]:
                    e["client_name"] = proj.get("client_name") or (
                        user.clients(proj.get("workspace_id") or wid).get(cid) if cid else None)
            elif not e["project_name"]:
                e["project_name"] = f"project #{pid}"
        if not e["tag_names"] and e["tag_ids"]:
            tmap = user.tags(wid)
            e["tag_names"] = [tmap.get(t, f"tag #{t}") for t in e["tag_ids"]]
        if e["task_id"] and not e["task_name"]:
            e["task_name"] = f"task #{e['task_id']}"


# ============================================================================
# Выгрузка по одному человеку
# ============================================================================
def fetch_person(abbr: str, token: str, start: dt.date, end: dt.date):
    user = TogglUser(abbr, token)
    me = user.load_me()
    tz_name = TIMEZONE_OVERRIDE or me.get("timezone") or "UTC"
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        log(f"   ⚠️ Неизвестный часовой пояс '{tz_name}', использую UTC")
        tz_name, tz = "UTC", ZoneInfo("UTC")

    start_local = dt.datetime.combine(start, dt.time.min, tz)
    end_local_excl = dt.datetime.combine(end + dt.timedelta(days=1), dt.time.min, tz)
    start_utc = start_local.astimezone(dt.timezone.utc)
    end_utc = end_local_excl.astimezone(dt.timezone.utc)

    source = SOURCE
    if source == "auto":
        border = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=TIMELINE_MAX_DAYS_BACK)
        source = "timeline" if start_utc >= border else "reports"

    entries = []
    if source == "timeline":
        try:
            entries = [_from_timeline(e) for e in user.timeline_entries(start_utc, end_utc)]
        except TimelineTooOld:
            log("   ℹ️ Период старше 3 месяцев для /me/time_entries — переключаюсь на Reports API")
            source = "reports"
    if source == "reports":
        wid, uid = me.get("default_workspace_id"), me.get("id")
        # запас ±1 день на случай, если Reports API считает даты в UTC; лишнее отфильтруем
        rows = user.reports_entries(wid, uid, start - dt.timedelta(days=1),
                                    end + dt.timedelta(days=1))
        entries = _from_reports(rows, wid)

    # дубликаты (стык при делении периода)
    entries = list({e["id"]: e for e in entries if e.get("id") is not None}.values())

    running = [e for e in entries if not e["stop"] or (e["seconds"] or 0) < 0]
    if running and not INCLUDE_RUNNING:
        entries = [e for e in entries if e not in running]

    resolve_names(user, entries)

    info = {
        "abbr": abbr, "name": me.get("fullname") or "", "email": me.get("email") or "",
        "tz_name": tz_name, "tz": tz, "source": source, "running": len(running),
        "requests": user.requests_made,
        "start_local": start_local, "end_local_excl": end_local_excl,
    }
    return entries, info


def to_dataframe(entries: list, info: dict) -> pd.DataFrame:
    cols = ["Abbreviation", "User", "Email", "Client", "Project", "Task", "Description",
            "Billable", "Start date", "Start time", "End date", "End time", "Duration",
            "Duration (h)", "Tags", "Time entry ID", "Workspace ID", "Timezone"]
    if not entries:
        return pd.DataFrame(columns=cols)

    df = pd.DataFrame(entries)
    tz = info["tz"]
    start = pd.to_datetime(df["start"], utc=True).dt.tz_convert(tz)
    stop = pd.to_datetime(df["stop"], utc=True, errors="coerce").dt.tz_convert(tz)
    now = pd.Timestamp.now(tz=tz)
    stop_filled = stop.fillna(now)
    seconds = pd.to_numeric(df["seconds"], errors="coerce")
    seconds = seconds.where(seconds >= 0, (stop_filled - start).dt.total_seconds()).round()

    # Строго по дате начала в локальном часовом поясе — как в отчётах Toggl
    mask = (start >= pd.Timestamp(info["start_local"])) & (start < pd.Timestamp(info["end_local_excl"]))
    df, start, stop, seconds = df[mask], start[mask], stop[mask], seconds[mask]

    def hms(s):
        s = int(s)
        return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"

    out = pd.DataFrame({
        "Abbreviation": info["abbr"],
        "User": info["name"],
        "Email": info["email"],
        "Client": df["client_name"].fillna(""),
        "Project": df["project_name"].fillna(""),
        "Task": df["task_name"].fillna(""),
        "Description": df["description"].fillna(""),
        "Billable": df["billable"].map(lambda b: "Yes" if b else "No"),
        "Start date": start.dt.date,
        "Start time": start.dt.strftime("%H:%M:%S"),
        "End date": stop.dt.date,
        "End time": stop.dt.strftime("%H:%M:%S"),
        "Duration": seconds.map(hms),
        "Duration (h)": (seconds / 3600).round(4),
        "Tags": df["tag_names"].map(lambda t: ", ".join(t) if isinstance(t, list) else ""),
        "Time entry ID": df["id"],
        "Workspace ID": df["workspace_id"],
        "Timezone": info["tz_name"],
    })
    out["_sort"] = start.dt.tz_localize(None)
    out = out.sort_values("_sort").drop(columns="_sort").reset_index(drop=True)
    return out[cols]


# ============================================================================
# Запись файлов
# ============================================================================
def _safe_path(path: Path) -> Path:
    """Если файл открыт в Excel (заблокирован) — пишем рядом с меткой времени."""
    try:
        if path.exists():
            with open(path, "a"):
                pass
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{dt.datetime.now():%H%M%S}{path.suffix}")
        log(f"   ⚠️ {path.name} открыт в другой программе — сохраняю как {alt.name}")
        return alt


def _format_sheet(ws, df: pd.DataFrame) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F4F4F")
    for i, col in enumerate(df.columns, start=1):
        sample = df[col].astype(str).head(500)
        width = max([len(str(col))] + [len(v) for v in sample]) if len(sample) else len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
        if col in ("Start date", "End date"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "yyyy-mm-dd"


def build_summaries(detail: pd.DataFrame):
    by_user = (detail.groupby(["Abbreviation", "User"], dropna=False)
               .agg(Entries=("Time entry ID", "count"), Hours=("Duration (h)", "sum"))
               .reset_index())
    by_project = (detail.groupby(["Abbreviation", "Client", "Project"], dropna=False)
                  .agg(Entries=("Time entry ID", "count"), Hours=("Duration (h)", "sum"))
                  .reset_index())
    by_user["Hours"] = by_user["Hours"].round(2)
    by_project["Hours"] = by_project["Hours"].round(2)
    return by_user, by_project


def write_outputs(detail: pd.DataFrame, run_log: pd.DataFrame, base_name: str,
                  out_dir: Path, per_person_sheets: bool) -> list:
    written = []
    if OUTPUT_FORMAT in ("xlsx", "both"):
        path = _safe_path(out_dir / f"{base_name}.xlsx")
        by_user, by_project = build_summaries(detail)
        sheets = {"Detailed": detail, "Summary by user": by_user,
                  "Summary by project": by_project, "Run log": run_log}
        if per_person_sheets:
            for abbr, part in detail.groupby("Abbreviation"):
                sheets[str(abbr)[:31]] = part
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for name, frame in sheets.items():
                frame.to_excel(writer, sheet_name=name, index=False)
                _format_sheet(writer.sheets[name], frame)
        written.append(path)
    if OUTPUT_FORMAT in ("csv", "both"):
        path = _safe_path(out_dir / f"{base_name}.csv")
        detail.to_csv(path, index=False, sep=CSV_SEP, encoding="utf-8-sig")
        written.append(path)
    return written


# ============================================================================
# Токены, аргументы, main
# ============================================================================
def load_tokens() -> dict:
    for cand in TOKENS_FILE_CANDIDATES:
        path = Path(cand.format(script_dir=SCRIPT_DIR))
        if path.exists():
            log(f"🔑 Токены: {path}")
            df = pd.read_excel(path, sheet_name=TOKENS_SHEET)
            tokens = {}
            for _, row in df.iterrows():
                abbr, tok = row.get(TOKENS_ABBR_COL), row.get(TOKENS_TOKEN_COL)
                if pd.notna(abbr) and pd.notna(tok) and str(tok).strip():
                    tokens[norm_abbr(abbr)] = str(tok).strip()
            return tokens
    raise FileNotFoundError("Не найден файл токенов. Проверьте TOKENS_FILE_CANDIDATES:\n  "
                            + "\n  ".join(TOKENS_FILE_CANDIDATES))


def parse_args():
    p = argparse.ArgumentParser(description="Детальный отчёт Toggl → Excel/CSV")
    p.add_argument("--start", help="YYYY-MM-DD, включительно")
    p.add_argument("--end", help="YYYY-MM-DD, включительно")
    p.add_argument("--month", help="YYYY-MM — весь месяц (вместо --start/--end)")
    p.add_argument("--people", nargs="+", help="аббревиатуры через пробел или ALL")
    p.add_argument("--format", choices=["xlsx", "csv", "both"])
    p.add_argument("--mode", choices=["one_file", "file_per_person"])
    p.add_argument("--source", choices=["auto", "timeline", "reports"])
    p.add_argument("--sheet-per-person", action="store_true")
    args, _ = p.parse_known_args()  # parse_known_args — чтобы не ломалось в Jupyter
    return args


def main():
    global OUTPUT_FORMAT, OUTPUT_MODE, SOURCE, SHEET_PER_PERSON
    args = parse_args()

    if args.month:
        y, m = map(int, args.month.split("-"))
        start = dt.date(y, m, 1)
        end = (dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1))
    else:
        start = dt.date.fromisoformat(args.start or START_DATE)
        end = dt.date.fromisoformat(args.end or END_DATE)
    if start > end:
        sys.exit(f"❌ Дата начала {start} позже даты конца {end}")

    OUTPUT_FORMAT = args.format or OUTPUT_FORMAT
    OUTPUT_MODE = args.mode or OUTPUT_MODE
    SOURCE = args.source or SOURCE
    SHEET_PER_PERSON = args.sheet_per_person or SHEET_PER_PERSON

    tokens = load_tokens()
    if args.people and [x.upper() for x in args.people] != ["ALL"]:
        people = [norm_abbr(x) for x in args.people]
    elif args.people or PEOPLE is None:
        people = list(tokens)
    else:
        people = [norm_abbr(x) for x in PEOPLE]

    out_dir = Path(OUTPUT_DIR) if OUTPUT_DIR else SCRIPT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    period = f"{start:%Y-%m-%d}_{end:%Y-%m-%d}"
    log(f"📅 Период: {start} … {end} | людей: {len(people)} | формат: {OUTPUT_FORMAT} "
        f"| режим: {OUTPUT_MODE} | источник: {SOURCE}")

    frames, log_rows, written = [], [], []
    for i, abbr in enumerate(people, 1):
        log(f"############ [{i}/{len(people)}] {abbr} ############")
        row = {"Abbreviation": abbr, "Status": "", "Source": "", "Entries": 0,
               "Hours": 0.0, "Skipped running": 0, "API requests": 0, "Message": ""}
        token = tokens.get(abbr)
        if not token:
            row.update(Status="ERROR", Message="Нет токена в файле токенов")
            log("   ❌ Нет токена — пропускаю")
            log_rows.append(row)
            continue
        try:
            entries, info = fetch_person(abbr, token, start, end)
            df = to_dataframe(entries, info)
            frames.append(df)
            row.update(Status="OK", Source=info["source"], Entries=len(df),
                       Hours=round(df["Duration (h)"].sum(), 2),
                       **{"Skipped running": info["running"] if not INCLUDE_RUNNING else 0,
                          "API requests": info["requests"]})
            log(f"   ✅ {len(df)} записей, {row['Hours']} ч (источник: {info['source']}, "
                f"запросов: {info['requests']})")
            if OUTPUT_MODE == "file_per_person":
                written += write_outputs(df, pd.DataFrame([row]), f"toggl_detailed_{abbr}_{period}",
                                         out_dir, per_person_sheets=False)
        except Exception as e:
            row.update(Status="ERROR", Message=str(e)[:500])
            log(f"   ❌ {e}")
        log_rows.append(row)

    run_log = pd.DataFrame(log_rows)
    if OUTPUT_MODE == "one_file":
        detail = pd.concat(frames, ignore_index=True) if frames else to_dataframe([], {})
        suffix = people[0] if len(people) == 1 else "ALL" if len(people) == len(tokens) else "team"
        written += write_outputs(detail, run_log, f"toggl_detailed_{suffix}_{period}",
                                 out_dir, per_person_sheets=SHEET_PER_PERSON)

    log("=" * 60)
    errors = run_log[run_log["Status"] != "OK"]
    log(f"Готово: OK {len(run_log) - len(errors)}, ошибок {len(errors)}")
    for _, r in errors.iterrows():
        log(f"   ❌ {r['Abbreviation']}: {r['Message']}")
    for p in written:
        log(f"💾 {p}")


if __name__ == "__main__":
    main()
